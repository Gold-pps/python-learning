"""四种格式解析层的单元测试。

样本全部在运行时动态生成（写进 tmp_path），不依赖仓库里的真实文件——
沿用第 14 周的原则：测试的清理责任交给 pytest，不写 finally。
"""

from __future__ import annotations

import pytest
from docx import Document
from fpdf import FPDF
from rag.parsers import parse_docx, parse_file, parse_md, parse_pdf, parse_txt


def test_parse_md_keeps_title_path(tmp_path):
    p = tmp_path / "t.md"
    p.write_text("# 标题一\naaa\n\n## 小标题\nbbb\nccc\n", encoding="utf-8")
    blocks = parse_md(p)
    assert len(blocks) == 2
    assert blocks[0].locator == "标题一"
    assert blocks[0].text == "aaa"
    assert blocks[1].locator == "标题一 > 小标题"
    assert blocks[1].text == "bbb\nccc"


def test_parse_txt_splits_on_blank_lines(tmp_path):
    p = tmp_path / "t.txt"
    p.write_text("第一段\nline2\n\n第二段\n", encoding="utf-8")
    blocks = parse_txt(p)
    assert len(blocks) == 2
    assert blocks[0].start == 1 and blocks[0].end == 2
    assert blocks[1].start == 4 and blocks[1].end == 4


def test_parse_pdf_extracts_per_page_text(tmp_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "page one")
    pdf.add_page()
    pdf.cell(0, 10, "page two")
    pdf_path = tmp_path / "t.pdf"
    pdf.output(str(pdf_path))

    blocks = parse_pdf(pdf_path)
    assert len(blocks) == 2
    assert (
        blocks[0].locator == "第 1 页" and blocks[0].start == 1 and blocks[0].end == 1
    )
    assert "page one" in blocks[0].text
    assert blocks[1].locator == "第 2 页"
    assert "page two" in blocks[1].text


def test_parse_docx_keeps_heading_path(tmp_path):
    doc = Document()
    doc.add_heading("标题一", level=1)
    doc.add_paragraph("aaa")
    doc.add_heading("小标题", level=2)
    doc.add_paragraph("bbb")
    docx_path = tmp_path / "t.docx"
    doc.save(str(docx_path))

    blocks = parse_docx(docx_path)
    assert len(blocks) == 2
    assert blocks[0].locator == "标题一"
    assert blocks[0].text == "aaa"
    assert blocks[1].locator == "标题一 > 小标题"
    assert blocks[1].text == "bbb"


def test_parse_file_dispatches_by_suffix(tmp_path):
    p = tmp_path / "t.md"
    p.write_text("# T\nbody\n", encoding="utf-8")
    blocks = parse_file(p)
    assert len(blocks) == 1
    assert blocks[0].text == "body"


def test_parse_file_rejects_unknown_suffix(tmp_path):
    p = tmp_path / "t.csv"
    p.write_text("a,b\n", encoding="utf-8")
    with pytest.raises(ValueError, match="不支持的格式"):
        parse_file(p)
